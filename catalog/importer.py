"""Excel (.xlsx) va CSV fayldan katalog yuklash.

Ustun nomlari o'zbek, rus yoki ingliz tilida bo'lishi mumkin.
Xatolar qator bo'yicha yig'iladi: bitta noto'g'ri qator butun yuklashni to'xtatmaydi.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.db import transaction
from openpyxl import Workbook, load_workbook

from .models import Product

COLUMN_ALIASES = {
    "name": {"nom", "nomi", "mahsulot", "mahsulot nomi", "название", "наименование", "товар", "name", "product"},
    "price": {"narx", "narxi", "цена", "price", "narxi som", "narxi sum"},
    "stock": {"qoldiq", "soni", "miqdor", "ombor", "остаток", "количество", "stock", "qty", "quantity"},
    "description": {"tavsif", "izoh", "ta'rif", "описание", "description"},
    "category": {"toifa", "kategoriya", "категория", "category"},
    "sku": {"kod", "artikul", "артикул", "sku", "code"},
}


class ImportErrorBase(Exception):
    """Fayl umuman o'qib bo'lmaydigan holatlar uchun."""


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def total_ok(self) -> int:
        return self.created + self.updated


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("’", "'").replace("`", "'"))


def _map_headers(headers: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, header in enumerate(headers):
        h = _norm_header(header)
        for field_name, aliases in COLUMN_ALIASES.items():
            if h in aliases and field_name not in mapping:
                mapping[field_name] = idx
    return mapping


def _parse_price(value) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError("narx bo'sh")
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
    else:
        cleaned = re.sub(r"[^\d.,]", "", str(value)).replace(",", ".")
        # "12.500" kabi minglik ajratgichni aniqlash
        if cleaned.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", cleaned):
            cleaned = cleaned.replace(".", "")
        try:
            number = Decimal(cleaned)
        except InvalidOperation as exc:
            raise ValueError(f"narx noto'g'ri: {value!r}") from exc
    if number < 0:
        raise ValueError("narx manfiy bo'lishi mumkin emas")
    return number.quantize(Decimal("1"))


def _parse_stock(value) -> int:
    if value is None or str(value).strip() == "":
        return 0
    try:
        number = int(Decimal(re.sub(r"[^\d.,-]", "", str(value)).replace(",", ".")))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"qoldiq noto'g'ri: {value!r}") from exc
    if number < 0:
        raise ValueError("qoldiq manfiy bo'lishi mumkin emas")
    return number


def _read_rows(filename: str, data: bytes) -> list[list]:
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl turli xatolar beradi
            raise ImportErrorBase("Excel faylni o'qib bo'lmadi. .xlsx formatini tekshiring.") from exc
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if lower.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportErrorBase("CSV fayl UTF-8 kodlashda bo'lishi kerak.") from exc
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    raise ImportErrorBase("Faqat .xlsx yoki .csv fayl qabul qilinadi.")


def import_catalog(shop, filename: str, data: bytes) -> ImportResult:
    if len(data) > settings.MAX_UPLOAD_BYTES:
        raise ImportErrorBase(f"Fayl juda katta (maksimum {settings.MAX_UPLOAD_BYTES // (1024 * 1024)} MB).")

    rows = _read_rows(filename, data)
    if not rows:
        raise ImportErrorBase("Fayl bo'sh.")

    mapping = _map_headers(rows[0])
    missing = [f for f in ("name", "price") if f not in mapping]
    if missing:
        raise ImportErrorBase(
            "Majburiy ustunlar topilmadi: nom va narx. Birinchi qatorda ustun nomlari bo'lishi kerak (shablonni yuklab oling)."
        )

    body = rows[1:]
    if len(body) > settings.MAX_CATALOG_ROWS:
        raise ImportErrorBase(f"Qatorlar soni {settings.MAX_CATALOG_ROWS} dan oshmasligi kerak.")

    def cell(row, key):
        idx = mapping.get(key)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    result = ImportResult()
    with transaction.atomic():
        for line_no, row in enumerate(body, start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            name = str(cell(row, "name") or "").strip()
            if not name:
                result.skipped += 1
                result.errors.append(f"{line_no}-qator: nom bo'sh")
                continue
            try:
                price = _parse_price(cell(row, "price"))
                stock = _parse_stock(cell(row, "stock"))
            except ValueError as exc:
                result.skipped += 1
                result.errors.append(f"{line_no}-qator ({name[:40]}): {exc}")
                continue

            sku = str(cell(row, "sku") or "").strip()[:64]
            defaults = {
                "name": name[:200],
                "price": price,
                "stock": stock,
                "description": str(cell(row, "description") or "").strip(),
                "category": str(cell(row, "category") or "").strip()[:120],
                "is_active": True,
            }
            existing = None
            if sku:
                existing = Product.objects.filter(shop=shop, sku=sku).first()
            if existing is None:
                existing = Product.objects.filter(shop=shop, name__iexact=name[:200]).first()
            if existing:
                for key, value in defaults.items():
                    setattr(existing, key, value)
                if sku:
                    existing.sku = sku
                existing.save()
                result.updated += 1
            else:
                Product.objects.create(shop=shop, sku=sku, **defaults)
                result.created += 1
    return result


def build_template_xlsx() -> bytes:
    """Egasi to'ldirishi uchun namunaviy Excel shablon."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Katalog"
    ws.append(["Nomi", "Narxi", "Qoldiq", "Toifa", "Tavsif", "Kod"])
    ws.append(["Erkaklar ko'ylagi (oq)", 180000, 12, "Kiyim", "Paxta, razmerlar: S, M, L, XL", "KY-001"])
    ws.append(["Ayollar sumkasi", 250000, 5, "Aksessuar", "Teri, qora rang", "SM-014"])
    ws.append(["Krossovka Run Pro", 420000, 8, "Poyabzal", "Razmerlar 40-45", "KR-220"])
    for col, width in zip("ABCDEF", (32, 14, 10, 16, 44, 12)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
